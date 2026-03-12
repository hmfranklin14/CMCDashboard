
import pandas as pd
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

# -------------------------
# 1. Load CSV
# -------------------------
df = pd.read_csv("/Users/harryfranklin/Downloads/Patients.csv")

# Optional: round numbers / clean formatting
df = df.round(2)

# -------------------------
# 2. Create Word document
# -------------------------
doc = Document()
doc.add_heading('Data Table', level=1)

# -------------------------
# 3. Create table
# -------------------------
table = doc.add_table(rows=1, cols=len(df.columns))
table.style = 'Table Grid'

# Header row
hdr_cells = table.rows[0].cells
for i, col in enumerate(df.columns):
    hdr_cells[i].text = col

# -------------------------
# 4. Fill table
# -------------------------
for _, row in df.iterrows():
    row_cells = table.add_row().cells
    for i, value in enumerate(row):
        row_cells[i].text = str(value)

# -------------------------
# 5. Adjust formatting
# -------------------------
for row in table.rows:
    for cell in row.cells:
        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                run.font.size = Pt(11)

# -------------------------
# 6. Save document
# -------------------------
doc.save("formatted_table.docx")